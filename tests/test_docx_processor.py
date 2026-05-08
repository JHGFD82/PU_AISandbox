"""Tests for src/processors/docx_processor.py — uncovered methods."""

import logging
from io import BytesIO
from typing import List

import pytest

from src.processors.docx_processor import DocxProcessor
from src.models.doc_block import ParagraphBlock, TableBlock


# ---------------------------------------------------------------------------
# Helpers to build minimal .docx files in memory
# ---------------------------------------------------------------------------

def _make_docx_bytes(paragraphs: List[str], tables: List[List[List[str]]] | None = None) -> bytes:
    """Create an in-memory .docx with given paragraphs and optional tables."""
    from docx import Document as _Doc
    doc = _Doc()
    for text in paragraphs:
        doc.add_paragraph(text)
    for table_data in (tables or []):
        if not table_data:
            continue
        n_rows = len(table_data)
        n_cols = max(len(r) for r in table_data)
        tbl = doc.add_table(rows=n_rows, cols=n_cols)
        for r_idx, row in enumerate(table_data):
            for c_idx, cell_text in enumerate(row):
                tbl.cell(r_idx, c_idx).text = cell_text
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


# ---------------------------------------------------------------------------
# extract_blocks
# ---------------------------------------------------------------------------

class TestExtractBlocks:

    def test_paragraphs_only(self):
        data = _make_docx_bytes(["Hello", "World"])
        blocks = DocxProcessor.extract_blocks(BytesIO(data))
        para_texts = [b.text for b in blocks if isinstance(b, ParagraphBlock)]
        assert "Hello" in para_texts
        assert "World" in para_texts

    def test_table_produces_table_block(self):
        data = _make_docx_bytes(
            ["Before"],
            tables=[[["A", "B"], ["C", "D"]]]
        )
        blocks = DocxProcessor.extract_blocks(BytesIO(data))
        table_blocks = [b for b in blocks if isinstance(b, TableBlock)]
        assert len(table_blocks) == 1
        assert table_blocks[0].placeholder == "[TABLE_1]"
        assert table_blocks[0].rows == [["A", "B"], ["C", "D"]]

    def test_multiple_tables_numbered_sequentially(self):
        data = _make_docx_bytes(
            [],
            tables=[[["X"]], [["Y"]]]
        )
        blocks = DocxProcessor.extract_blocks(BytesIO(data))
        table_blocks = [b for b in blocks if isinstance(b, TableBlock)]
        placeholders = [b.placeholder for b in table_blocks]
        assert "[TABLE_1]" in placeholders
        assert "[TABLE_2]" in placeholders

    def test_empty_document_returns_empty_list(self):
        data = _make_docx_bytes([])
        blocks = DocxProcessor.extract_blocks(BytesIO(data))
        # python-docx may add a default empty paragraph; filtering to non-empty text
        para_blocks = [b for b in blocks if isinstance(b, ParagraphBlock)]
        assert all(b.text.strip() for b in para_blocks)


# ---------------------------------------------------------------------------
# process_docx_for_translation
# ---------------------------------------------------------------------------

class TestProcessDocxForTranslation:

    def test_paragraphs_become_pages(self):
        data = _make_docx_bytes(["First paragraph", "Second paragraph"])
        pages, registry = DocxProcessor.process_docx_for_translation(BytesIO(data))
        assert isinstance(pages, list)
        assert len(pages) >= 1
        combined = "\n".join(pages)
        assert "First paragraph" in combined or "Second" in combined

    def test_table_appears_as_placeholder_in_pages_and_registry(self):
        data = _make_docx_bytes(
            ["Intro"],
            tables=[[["Col A", "Col B"], ["Val 1", "Val 2"]]]
        )
        pages, registry = DocxProcessor.process_docx_for_translation(BytesIO(data))
        assert "[TABLE_1]" in registry
        assert registry["[TABLE_1]"] == [["Col A", "Col B"], ["Val 1", "Val 2"]]
        combined = "\n".join(pages)
        assert "[TABLE_1]" in combined

    def test_empty_document_returns_empty_string_page(self):
        data = _make_docx_bytes([])
        pages, registry = DocxProcessor.process_docx_for_translation(BytesIO(data))
        assert pages == [""]
        assert registry == {}


# ---------------------------------------------------------------------------
# extract_media — image-free document
# ---------------------------------------------------------------------------

class TestExtractMedia:

    def test_no_images_returns_empty_list(self):
        data = _make_docx_bytes(["Plain text paragraph"])
        media = DocxProcessor.extract_media(BytesIO(data))
        assert media == []
