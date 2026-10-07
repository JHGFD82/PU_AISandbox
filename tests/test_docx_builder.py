"""Tests for src/output/docx_builder.py: saving results as a Word document.

Each test opens the document that was written and looks for what should be
in it — a picture, a table, the text — rather than only checking that saving
did not fail. Whenever a Word document cannot be made, the text must still
reach the person, as a plain .txt file beside where the .docx would have gone.
"""

import builtins
import io
import logging
import struct
import zlib
from unittest.mock import MagicMock, patch

from docx import Document

from src.models.embedded_media import EmbeddedMedia
from src.output import docx_builder
from src.output.docx_builder import _apply_docx_table_borders, _fallback_to_text, save_to_docx


def _png() -> bytes:
    """A 1×1 red PNG, built by hand so the test needs nothing else to make one."""
    def chunk(tag, data):
        c = struct.pack(">I", len(data)) + tag + data
        return c + struct.pack(">I", zlib.crc32(c[4:]) & 0xFFFFFFFF)
    header = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    return (b"\x89PNG\r\n\x1a\n"
            + chunk(b"IHDR", header)
            + chunk(b"IDAT", zlib.compress(b"\x00\xff\x00\x00"))
            + chunk(b"IEND", b""))


def _webp() -> bytes:
    """A small WebP image: a format Word will not take, but Pillow can convert."""
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (4, 4), "red").save(buf, format="WEBP")
    return buf.getvalue()


def _media(data: bytes, content_type: str = "image/png") -> EmbeddedMedia:
    return EmbeddedMedia(
        data=data,
        content_type=content_type,
        position_fraction=0.5,
        width_emu=914400,
        height_emu=914400,
    )


def _texts(path) -> list[str]:
    return [p.text for p in Document(str(path)).paragraphs if p.text]


class TestTableBorders:
    def test_a_failure_is_logged_rather_than_raised(self, caplog):
        bad_table = MagicMock()
        bad_table._tbl.tblPr.append.side_effect = RuntimeError("xml error")
        with caplog.at_level(logging.WARNING):
            _apply_docx_table_borders(bad_table)
        assert any("Could not apply table borders" in r.message for r in caplog.records)


class TestFallingBackToText:
    """When no Word document can be made, the text is saved as .txt instead."""

    def test_without_python_docx(self, tmp_path, monkeypatch):
        real_import = builtins.__import__

        def no_docx(name, *args, **kwargs):
            if name == "docx":
                raise ImportError("No module named 'docx'")
            return real_import(name, *args, **kwargs)

        monkeypatch.setattr(builtins, "__import__", no_docx)
        save_to_docx("Some content", str(tmp_path / "result.docx"), label="Test")

        assert not (tmp_path / "result.docx").exists()
        assert "Some content" in (tmp_path / "result.txt").read_text(encoding="utf-8")

    def test_when_making_the_document_fails(self, tmp_path):
        with patch("docx.Document", side_effect=RuntimeError("disk full")):
            save_to_docx("Content here", str(tmp_path / "output.docx"), label="Translation")

        assert not (tmp_path / "output.docx").exists()
        assert "Content here" in (tmp_path / "output.txt").read_text(encoding="utf-8")

    def test_when_there_is_nothing_to_put_in_it(self, tmp_path):
        save_to_docx("", str(tmp_path / "empty.docx"), label="Translation")

        assert not (tmp_path / "empty.docx").exists()
        assert (tmp_path / "empty.txt").exists()

    def test_the_fallback_swaps_the_extension(self, tmp_path):
        _fallback_to_text("Hello world", str(tmp_path / "output.docx"), "Test")
        assert "Hello world" in (tmp_path / "output.txt").read_text(encoding="utf-8")


class TestPictures:
    def test_a_picture_is_placed_among_the_text(self, tmp_path):
        out = tmp_path / "with_image.docx"
        save_to_docx("Paragraph one\n\nParagraph two", str(out),
                     media=[_media(_png())], label="Translation")

        assert len(Document(str(out)).inline_shapes) == 1
        assert _texts(out) == ["Paragraph one", "Paragraph two"]

    def test_a_format_word_will_not_take_is_converted_first(self, tmp_path):
        out = tmp_path / "webp.docx"
        save_to_docx("One\n\nTwo", str(out),
                     media=[_media(_webp(), "image/webp")], label="Translation")

        assert len(Document(str(out)).inline_shapes) == 1

    def test_an_unreadable_picture_is_left_out_and_the_text_kept(self, tmp_path, caplog):
        out = tmp_path / "bad_image.docx"
        with caplog.at_level(logging.WARNING):
            save_to_docx("Some text", str(out),
                         media=[_media(b"not an image")], label="Translation")

        assert any("Could not insert image" in r.message for r in caplog.records)
        assert len(Document(str(out)).inline_shapes) == 0
        assert _texts(out) == ["Some text"]


class TestTables:
    def test_a_placeholder_becomes_a_word_table(self, tmp_path):
        out = tmp_path / "with_table.docx"
        registry = {"[TABLE_1]": [["Header A", "Header B"], ["Row 1", "Row 2"]]}
        save_to_docx("[TABLE_1]\n\nSome other paragraph", str(out),
                     table_registry=registry, label="Test")

        tables = Document(str(out)).tables
        assert len(tables) == 1
        assert [c.text for c in tables[0].rows[1].cells] == ["Row 1", "Row 2"]
        assert "[TABLE_1]" not in _texts(out)

    def test_a_table_that_cannot_be_built_is_written_as_its_placeholder(
        self, tmp_path, monkeypatch, caplog
    ):
        """The placeholder text stays visible, so nothing vanishes without a trace."""
        def fail(_table):
            raise RuntimeError("table error")

        monkeypatch.setattr(docx_builder, "_apply_docx_table_borders", fail)
        out = tmp_path / "tbl_fail.docx"
        with caplog.at_level(logging.WARNING):
            save_to_docx("[TABLE_1]\n\nExtra paragraph", str(out),
                         table_registry={"[TABLE_1]": [["A", "B"]]}, label="Test")

        assert any("Could not insert table" in r.message for r in caplog.records)
        assert _texts(out) == ["[TABLE_1]", "Extra paragraph"]
