"""Small things many tests build: a picture, a Word document, a sandbox with nothing real behind it.

Kept here so each is written once. Core's tests and the plugins' tests both
import from this module, as ``from tests.helpers import png``.
"""

import struct
import zlib
from io import BytesIO
from unittest.mock import MagicMock


WHITE = (255, 255, 255)
RED = (255, 0, 0)


def png(width: int = 1, height: int = 1, colour: tuple[int, int, int] = WHITE) -> bytes:
    """A valid PNG of one solid colour, built by hand so no imaging library is needed."""
    def chunk(tag: bytes, data: bytes) -> bytes:
        c = struct.pack(">I", len(data)) + tag + data
        return c + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)

    header = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    # Each row starts with a filter byte of 0, then three bytes a pixel.
    row = b"\x00" + bytes(colour) * width
    return (b"\x89PNG\r\n\x1a\n"
            + chunk(b"IHDR", header)
            + chunk(b"IDAT", zlib.compress(row * height))
            + chunk(b"IEND", b""))


def docx_bytes(paragraphs: list[str], tables: list[list[list[str]]] | None = None) -> bytes:
    """A real Word document, as bytes, holding *paragraphs* and then any *tables*.

    Each table is a list of rows, each row a list of cell texts; rows may be of
    different lengths, and the table is as wide as its longest row.
    """
    from docx import Document

    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    for rows in tables or []:
        if not rows:
            continue
        table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
        for r, row in enumerate(rows):
            for c, text in enumerate(row):
                table.cell(r, c).text = text
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def bare_processor(monkeypatch, *services: str):
    """A SandboxProcessor with no API key, no usage file and no model catalog behind it.

    Built without running its own set-up, so nothing real is reached. Its file
    readers and output are stand-ins that record what they are asked, and so is
    each of *services*, named as the attributes a plugin's code expects to find
    on it, such as ``"translation_service"``.
    """
    # Imported here, not at the top: importing SandboxProcessor puts the class
    # together from whichever plugins have registered by then, and a module
    # imported early would fix it before they had.
    from src.runtime.sandbox_processor import SandboxProcessor

    monkeypatch.setattr("src.runtime.sandbox_processor.get_api_key",
                        lambda name: ("fake-key", "Professor Fake"))
    for name in ("resolve_model", "maybe_sync_model_pricing",
                 "get_model_system_role", "get_model_max_completion_tokens"):
        monkeypatch.setattr(f"src.services.base_service.{name}",
                            MagicMock(return_value="gpt-4o"), raising=False)
    monkeypatch.setattr("src.tracking.token_tracker.TokenTracker.__init__",
                        lambda self, professor: None)

    proc = SandboxProcessor.__new__(SandboxProcessor)
    proc.professor_name = "fake"
    proc.professor_display_name = "Professor Fake"
    proc.token_tracker = MagicMock()
    proc.token_tracker.usage_data = {"total_usage": {"total_tokens": 0, "total_cost": 0.0}}
    proc.image_processor = MagicMock()
    proc.pdf_processor = MagicMock()
    proc.file_output = MagicMock()
    for name in services:
        setattr(proc, name, MagicMock())
    return proc
