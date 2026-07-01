"""Tests for plugins/translation/src/processors/docx_translation.py (registered as src.processors.docx_translation)."""

import sys
from io import BytesIO
from typing import List

process_docx_for_translation = sys.modules["src.processors.docx_translation"].process_docx_for_translation


# ---------------------------------------------------------------------------
# Helpers to build minimal .docx files in memory
# ---------------------------------------------------------------------------

def _make_docx_bytes(paragraphs: List[str], tables: List[List[List[str]]] | None = None) -> bytes:
    """Create an in-memory .docx with given paragraphs and optional tables."""
    from docx import Document as _Doc
    doc = _Doc()
    for text in paragraphs:
        doc.add_paragraph(text)
    for table_data in (tables or []):
        if not table_data:
            continue
        n_rows = len(table_data)
        n_cols = max(len(r) for r in table_data)
        tbl = doc.add_table(rows=n_rows, cols=n_cols)
        for r_idx, row in enumerate(table_data):
            for c_idx, cell_text in enumerate(row):
                tbl.cell(r_idx, c_idx).text = cell_text
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


# ---------------------------------------------------------------------------
# process_docx_for_translation
# ---------------------------------------------------------------------------

class TestProcessDocxForTranslation:

    def test_paragraphs_become_pages(self):
        data = _make_docx_bytes(["First paragraph", "Second paragraph"])
        pages, registry = process_docx_for_translation(BytesIO(data))
        assert isinstance(pages, list)
        assert len(pages) >= 1
        combined = "\n".join(pages)
        assert "First paragraph" in combined or "Second" in combined

    def test_table_appears_as_placeholder_in_pages_and_registry(self):
        data = _make_docx_bytes(
            ["Intro"],
            tables=[[["Col A", "Col B"], ["Val 1", "Val 2"]]]
        )
        pages, registry = process_docx_for_translation(BytesIO(data))
        assert "[TABLE_1]" in registry
        assert registry["[TABLE_1]"] == [["Col A", "Col B"], ["Val 1", "Val 2"]]
        combined = "\n".join(pages)
        assert "[TABLE_1]" in combined

    def test_empty_document_returns_empty_string_page(self):
        data = _make_docx_bytes([])
        pages, registry = process_docx_for_translation(BytesIO(data))
        assert pages == [""]
        assert registry == {}
