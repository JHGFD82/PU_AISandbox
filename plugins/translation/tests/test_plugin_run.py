"""Tests for plugins/translation/plugin.py's run() and _execute_translate(): what ``translate`` does when typed.

The sandbox is a stand-in that records what it is asked to do, so these tests
check which flag combinations are refused, what a dry run previews, and which
work each input leads to — not the translating itself, which is tested with
the services and runtime handler.
"""

import argparse
import json
import os
from unittest.mock import MagicMock

import pytest

import plugins.translation.plugin as plugin_mod
from src.errors import CLIError
from tests.helpers import docx_bytes


def _make_sandbox():
    fake = MagicMock()
    fake._sampling_kwargs.return_value = {}
    fake._resolve_output_path.return_value = "/out/result.docx"
    fake._collect_notes.return_value = ("system note", "user note")
    fake._collect_multiline.return_value = ""
    fake.translation_service.build_prompts.return_value = ("text system", "text user")
    fake.translation_service._get_model.return_value = "gpt-4o"
    fake.translation_service.variant_notes = []
    fake.image_translation_service.build_prompts.return_value = ("image system", "image user")
    fake.image_translation_service._get_model.return_value = "gpt-4o-vision"
    fake.image_processor.is_image_file.return_value = False
    return fake


@pytest.fixture
def sandbox():
    return _make_sandbox()


def _args(**flags):
    defaults = dict(input_file=None, custom_text=False, output_file=None, auto_save=False,
                    progressive_save=False, preserve_media=False, scanned=False, notes=False,
                    dry_run=False, preserve_tables=False, toc=False, abstract=False,
                    workers=1, spread=False, page_nums=None, custom_font=None, font_size=None)
    return argparse.Namespace(**(defaults | flags))


def _translate(sandbox, **flags):
    plugin_mod._execute_translate(sandbox, _args(**flags), "Japanese", "English")


class TestRun:
    @pytest.fixture
    def built(self, monkeypatch, sandbox):
        constructor = MagicMock(return_value=sandbox)
        monkeypatch.setattr("src.runtime.sandbox_processor.SandboxProcessor", constructor)
        return constructor

    def test_the_sandbox_is_built_for_the_person_with_the_settings_given(self, built, sandbox):
        args = _args(language_code=("en", "en"), custom_text=True, command="translate")
        plugin_mod.plugin.run(args, "heller", "gpt-4o", 0.2, 0.9, 500)
        built.assert_called_once_with("heller", model="gpt-4o", temperature=0.2, top_p=0.9, max_tokens=500)
        sandbox.translate_custom_text.assert_called_once()

    def test_a_single_language_is_not_a_pair(self, built):
        with pytest.raises(CLIError, match="language pair"):
            plugin_mod.plugin.run(_args(language_code="en"), "heller", None, None, None, None)

    def test_guidance_from_the_plugin_owning_the_target_language_is_added(self, built, sandbox):
        args = _args(language_code=("en", "en"), custom_text=True, _peer_guidance=["use British spelling"])
        plugin_mod.plugin.run(args, "heller", None, None, None, None)
        assert sandbox.translation_service.variant_notes == ["use British spelling"]

    def test_english_asks_for_no_guidance_of_its_own(self):
        assert plugin_mod.plugin.get_peer_guidance("en") is None


class TestScannedNeedsAPdf:
    @pytest.mark.parametrize("flags,reason", [
        (dict(), "requires a file input"),
        (dict(input_file="doc.pdf", custom_text=True), "custom text"),
        (dict(input_file="photo.png"), "only valid for PDF files"),
        (dict(input_file="doc.pdf", preserve_media=True), "--preserve-media"),
    ], ids=["no-input", "pasted-text", "not-a-pdf", "with-preserve-media"])
    def test_what_it_cannot_be_combined_with(self, sandbox, flags, reason):
        with pytest.raises(CLIError, match=reason):
            _translate(sandbox, scanned=True, **flags)


class TestPreserveMediaNeedsAWordDocumentOut:
    @pytest.mark.parametrize("flags,reason", [
        (dict(input_file="a.docx", output_file="b.docx", progressive_save=True), "--progressive-save"),
        (dict(custom_text=True), "pasted text contains no embedded media"),
        (dict(), "without a file input"),
        (dict(input_file="photo.jpg", output_file="b.docx"), "images have no embedded media"),
        (dict(input_file="notes.txt", output_file="b.docx"), "'.txt' files"),
        (dict(input_file="a.docx", auto_save=True), "auto-save produces a .txt"),
        (dict(input_file="a.docx"), "without a .docx output"),
        (dict(input_file="a.docx", output_file="b.txt"), "cannot embed images"),
        (dict(input_file="a.docx", output_file="b.pdf"), "does not yet support PDF output"),
        (dict(input_file="a.docx", output_file="b.md"), "got '.md'"),
    ], ids=["progressive", "pasted", "no-input", "image-in", "txt-in", "auto-save",
            "no-output", "txt-out", "pdf-out", "md-out"])
    def test_what_it_refuses(self, sandbox, flags, reason):
        with pytest.raises(CLIError, match=reason):
            _translate(sandbox, preserve_media=True, **flags)

    @pytest.mark.parametrize("source", ["a.docx", "a.pdf"])
    def test_a_word_document_or_pdf_into_a_word_document_goes_ahead(self, sandbox, source):
        _translate(sandbox, preserve_media=True, input_file=source, output_file="b.docx")
        opts = sandbox.translate_document.call_args.args[5]
        assert opts.preserve_media is True


class TestNotes:
    def test_for_a_picture_the_picture_prompts_are_shown(self, sandbox, tmp_path):
        picture = tmp_path / "page.png"
        picture.write_bytes(b"x")
        sandbox.image_processor.is_image_file.return_value = True
        _translate(sandbox, notes=True, input_file=str(picture))
        sandbox._collect_notes.assert_called_once_with("image system", "image user")

    def test_for_a_document_the_text_prompts_are_shown(self, sandbox, tmp_path):
        letter = tmp_path / "letter.txt"
        letter.write_text("x")
        _translate(sandbox, notes=True, input_file=str(letter))
        assert "[Japanese document text]" in sandbox.translation_service.build_prompts.call_args.args[0]
        sandbox._collect_notes.assert_called_once_with("text system", "text user")

    def test_for_pasted_text_the_text_prompts_are_shown(self, sandbox):
        _translate(sandbox, notes=True, custom_text=True)
        assert "[Japanese custom text]" in sandbox.translation_service.build_prompts.call_args.args[0]

    def test_the_notes_reach_both_services(self, sandbox):
        """Which one runs depends on what turns up in the input, so both are told."""
        _translate(sandbox, notes=True, custom_text=True)
        for service in (sandbox.translation_service, sandbox.image_translation_service):
            assert (service.system_note, service.user_note) == ("system note", "user note")

    def test_notes_typed_on_the_command_line_reach_both_services(self, sandbox):
        _translate(sandbox, custom_text=True)
        applied = [c.args[0] for c in sandbox._apply_inline_notes.call_args_list]
        assert applied == [sandbox.translation_service, sandbox.image_translation_service]


class TestOptionsPassedOn:
    def test_preserving_tables_is_turned_on_in_both_services(self, sandbox):
        _translate(sandbox, custom_text=True, preserve_tables=True)
        assert sandbox.translation_service.tables is True
        assert sandbox.image_translation_service.tables is True

    def test_a_table_of_contents_is_asked_for(self, sandbox):
        _translate(sandbox, custom_text=True, toc=True)
        assert sandbox.translation_service.toc is True


class TestDryRun:
    """Shows what would be sent for the first page, and sends nothing."""

    def _shown(self, sandbox):
        model, system, user = sandbox._dry_run_display.call_args.args[:3]
        return model, system, user

    def _first_page_sent(self, sandbox):
        return sandbox.translation_service.build_prompts.call_args.args[0]

    def test_a_picture_previews_the_picture_prompts(self, sandbox, tmp_path):
        sandbox._detect_and_validate_file.return_value = "image"
        _translate(sandbox, dry_run=True, input_file=str(tmp_path / "page.png"))
        assert self._shown(sandbox) == ("gpt-4o-vision", "image system", "image user")
        sandbox.translate_document.assert_not_called()

    def test_a_scanned_pdf_previews_the_picture_prompts(self, sandbox, tmp_path):
        sandbox._detect_and_validate_file.return_value = "pdf"
        _translate(sandbox, dry_run=True, scanned=True, input_file=str(tmp_path / "scan.pdf"))
        assert self._shown(sandbox)[1:] == ("image system", "image user")
        assert "Scanned PDF" in sandbox._dry_run_display.call_args.kwargs["note"]

    def test_a_pdf_previews_its_first_page(self, sandbox, tmp_path):
        pdf = tmp_path / "doc.pdf"
        pdf.write_bytes(b"%PDF")
        sandbox._detect_and_validate_file.return_value = "pdf"
        sandbox.pdf_processor.process_pdf.return_value = iter(["page one"])
        sandbox.pdf_processor.process_page.return_value = "First page of the PDF."
        _translate(sandbox, dry_run=True, input_file=str(pdf))
        assert "First page of the PDF." in self._first_page_sent(sandbox)
        assert self._shown(sandbox) == ("gpt-4o", "text system", "text user")

    def test_a_pdf_with_no_text_says_so(self, sandbox, tmp_path):
        pdf = tmp_path / "doc.pdf"
        pdf.write_bytes(b"%PDF")
        sandbox._detect_and_validate_file.return_value = "pdf"
        sandbox.pdf_processor.process_pdf.return_value = iter([])
        _translate(sandbox, dry_run=True, input_file=str(pdf))
        assert "[no text found in PDF]" in self._first_page_sent(sandbox)

    @pytest.mark.parametrize("kind,name,content,expected", [
        ("docx", "letter.docx", docx_bytes(["Dear colleague,"]), "Dear colleague,"),
        ("txt", "letter.txt", "Plain words.".encode(), "Plain words."),
        ("json", "data.json", json.dumps({"title": "A record"}).encode(), "A record"),
        ("markdown", "notes.md", "# Heading\n\nSome notes.".encode(), "Some notes."),
    ])
    def test_a_document_previews_its_first_page(self, sandbox, tmp_path, kind, name, content, expected):
        path = tmp_path / name
        path.write_bytes(content)
        sandbox._detect_and_validate_file.return_value = kind
        _translate(sandbox, dry_run=True, input_file=str(path))
        assert expected in self._first_page_sent(sandbox)

    def test_a_spreadsheet_previews_its_first_page(self, sandbox, tmp_path):
        from openpyxl import Workbook

        book = Workbook()
        book.active.append(["Name", "Year"])
        book.active.append(["Hokusai", "1831"])
        path = tmp_path / "sheet.xlsx"
        book.save(path)
        sandbox._detect_and_validate_file.return_value = "excel"
        _translate(sandbox, dry_run=True, input_file=str(path))
        assert "Hokusai" in self._first_page_sent(sandbox)

    def test_a_type_it_has_no_reader_for_previews_a_placeholder(self, sandbox, tmp_path):
        sandbox._detect_and_validate_file.return_value = "something-else"
        _translate(sandbox, dry_run=True, input_file=str(tmp_path / "x.bin"))
        assert "[Japanese text to translate]" in self._first_page_sent(sandbox)

    def test_pasted_text_is_previewed(self, sandbox):
        sandbox._collect_multiline.return_value = "Pasted words."
        _translate(sandbox, dry_run=True, custom_text=True)
        assert "Pasted words." in self._first_page_sent(sandbox)
        sandbox.translate_custom_text.assert_not_called()

    def test_nothing_pasted_previews_a_placeholder(self, sandbox):
        _translate(sandbox, dry_run=True, custom_text=True)
        assert "[Japanese text to translate]" in self._first_page_sent(sandbox)

    def test_no_input_previews_a_placeholder(self, sandbox):
        _translate(sandbox, dry_run=True)
        assert "[Japanese text to translate]" in self._first_page_sent(sandbox)

    def test_an_abstract_is_previewed_as_the_context(self, sandbox):
        sandbox._collect_multiline.return_value = "The abstract."
        _translate(sandbox, dry_run=True, abstract=True)
        assert "The abstract." in self._first_page_sent(sandbox)
        assert sandbox.translation_service.build_prompts.call_args.kwargs["context_type"] == "abstract"

    @pytest.mark.parametrize("flags,expected", [
        (dict(output_file="out.pdf"), "pdf"),
        (dict(output_file="out.docx"), "docx"),
        (dict(output_file="out.xlsx"), "xlsx"),
        (dict(output_file="out.odt"), "file"),
        (dict(output_file="out"), "file"),
        (dict(auto_save=True), "txt"),
        (dict(), "console"),
    ])
    def test_the_prompt_is_told_where_the_result_will_go(self, sandbox, flags, expected):
        _translate(sandbox, dry_run=True, **flags)
        assert sandbox.translation_service.build_prompts.call_args.kwargs["output_format"] == expected


class TestTranslating:
    def test_pasted_text_is_translated(self, sandbox):
        _translate(sandbox, custom_text=True)
        source, target, abstract, opts = sandbox.translate_custom_text.call_args.args
        assert (source, target, abstract) == ("Japanese", "English", None)
        assert opts.output_file == "/out/result.docx"

    def test_a_document_is_translated_with_the_pages_and_options_asked_for(self, sandbox):
        _translate(sandbox, input_file="letter.pdf", page_nums="1-3", workers=4,
                   scanned=True, auto_save=True, font_size=12, custom_font="Garamond")
        call = sandbox.translate_document.call_args
        assert call.args[:4] == ("letter.pdf", "Japanese", "English", "1-3")
        opts = call.args[5]
        assert (opts.auto_save, opts.font_size, opts.custom_font) == (True, 12, "Garamond")
        assert call.kwargs == {"workers": 4, "spread": False, "scanned": True}

    def test_a_folder_of_pictures_is_translated(self, sandbox, tmp_path):
        _translate(sandbox, input_file=str(tmp_path), workers=3, spread=True)
        call = sandbox.process_image_translation_folder.call_args
        assert call.args[:3] == (os.path.abspath(str(tmp_path)), "Japanese", "English")
        assert call.kwargs == {"workers": 3, "spread": True}

    def test_an_abstract_is_asked_for_and_passed_on(self, sandbox):
        sandbox._collect_multiline.return_value = "The abstract."
        _translate(sandbox, custom_text=True, abstract=True)
        assert sandbox.translate_custom_text.call_args.args[2] == "The abstract."

    def test_no_input_says_how_to_give_one(self, sandbox):
        with pytest.raises(CLIError, match="Use -i for file input or -c"):
            _translate(sandbox)


class TestFromTheWebForm:
    def test_fewer_than_one_worker_is_refused(self, monkeypatch, sandbox, tmp_path):
        monkeypatch.setattr("src.runtime.sandbox_processor.SandboxProcessor", MagicMock(return_value=sandbox))
        letter = tmp_path / "letter.txt"
        letter.write_text("x")
        with pytest.raises(CLIError, match="at least 1"):
            plugin_mod.plugin.run_ui_action(
                fields={"source_language": "en", "target_language": "en",
                        "file_path": str(letter), "workers": "-2"},
                professor="heller", model=None, on_progress=None,
                output_dir=str(tmp_path / "job_output"),
            )
